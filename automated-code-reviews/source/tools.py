from git import Repo
from state import GraphState
import os, json, datetime
from google.genai.types import CreateCachedContentConfig, GenerateContentConfig
from google import genai
from langchain_google_vertexai import VertexAI, ChatVertexAI
from docx import Document
from google.cloud import storage


# Environment variables
org_name = os.environ.get("org_name")
local_directory = os.environ.get("local_directory")
pat = os.environ.get("PAT")
project_id = os.environ["project_id"]
location = os.environ["location"]
model_id = os.environ.get("model_id")
gcs_bucket = os.environ["gcs_bucket"]
local_directory = os.environ.get("local_directory")
# End of Environment variables


# Initialize the GenAI client
client = genai.Client(
      vertexai=True, project=project_id,
      location=location,)


# Initialize the Storage client
client = storage.Client()
bucket = client.bucket(gcs_bucket)


# Node
def clone_repository(state: GraphState):
    """
    Clones a Git repository to a specified destination.
    """

    git_org_name = state.get("git_org_name")
    git_repository_name = state.get("git_repository_name")


    repo_url = f"https://{git_org_name}:{pat}@github.com/{git_org_name}/{git_repository_name}.git"
    destination_path = f"{local_directory}/{git_repository_name}"


    try:
        Repo.clone_from(repo_url, destination_path)
        print(f"Repository cloned successfully to {destination_path}")
    except Exception as e:
        print(f"An error occurred: {e}")

    return { "cloned_repository_path" : destination_path }


# Node
def read_repository_contents(state: GraphState):

    """ Reads the repository and copy all contents to a single file. """

    repo_path = state.get("cloned_repository_path")

    file_contents = {}

    for root, _, files in os.walk(repo_path):
        for file in files:
            # Exclude .git folder and its contents
            if ".git" not in root:
                file_path = os.path.join(root, file)
                try:
                    with open(file_path, "r", encoding="utf-8") as f:
                        file_contents[file_path] = f.read()
                except Exception as e:  # Catch any unexpected errors
                        print(f"An unexpected error occurred: {e}")

    output_file = "{}/complete_code.text".format(local_directory)

    if file_contents:
        with open(output_file, "w", encoding="utf-8") as outfile:

            for file_path, content in file_contents.items():
                outfile.write(f"<file path={file_path}>")
                outfile.write(f"{content}")
                outfile.write("</file>")
                outfile.write("\n \n")

        outfile.close()


    size_in_bytes = os.path.getsize(output_file)

    print(f"File Size: {size_in_bytes/1024/1024} MB")

    return { "consolidated_code_file_name" : output_file }


# Node
def caching_contents(state: GraphState):

    """ Cache the contents of the consolidated code file in Gemini Cache for multiple use. """


    file_name = state.get("consolidated_code_file_name")

    with open(file_name, "r", encoding="utf-8") as f:
        fullcode_as_string = f.read()

    contents = [ fullcode_as_string ]

    print(contents)

    system_instruction = """
        You are an software engineer & code reviewer.
        Stick to provided source code and never make up things.
        Check the following code base for the entire repository and answer questions.
    """

    # Initialize the GenAI client
    client = genai.Client(
        vertexai=True, project=project_id,
        location=location,)

    content_cache = client.caches.create(
        model=model_id,
        config=CreateCachedContentConfig(
            contents=contents,
            system_instruction=system_instruction,
            display_name="atos-prj-genai-cache",
            ttl="86400s",
        ),
    )

    print(content_cache.name)
    
    return { "cache_name" : content_cache.name }


def ask_question(cache_name, question):

    # Initialize the GenAI client
    client = genai.Client(
        vertexai=True, project=project_id,
        location=location,)

    response = client.models.generate_content(
        model=model_id,
        contents=question,
            config=GenerateContentConfig(
                cached_content=cache_name,
            ),
    )

    return response.text


# Node
def prepare_code_review(state: GraphState):

    """ Based on the questions it does the code reviews and notes in a docx file. """

    cache_name = state.get("cache_name")
    repository_name = state.get("git_repository_name")
    release_name = state.get("git_release_name")

    # Load questions from JSON file
    with open('questionnaires.json', 'r') as f:
        questionnaires = json.load(f)

    date_raw = datetime.datetime.now()

    document = Document()
    document.add_heading("Code Review Report", 0)

    file_name  = "{}_{}_{}_{}_{}.docx".format(repository_name, release_name,
                     date_raw.year,  date_raw.month,  date_raw.day)

    absolute_file_name = f'{local_directory}/{file_name}'

    document.add_heading('Coding Reviews', level=1)

    # Example usage: Loop through coding best practices questions
    for question in questionnaires['coding_best_practices']:
        response = ask_question(cache_name, question)        
        document.add_paragraph(f'Question: {question}')
        document.add_paragraph(f'Answer: {response}')
        
    document.add_page_break()

    document.add_heading('Security Best Practices', level=1)

    # Example usage: Loop through security best practices questions
    for question in questionnaires['security_best_practices']:
        response = ask_question(cache_name, question)
        document.add_paragraph(f'Question: {question}')
        document.add_paragraph(f'Answer: {response}')
        
    document.add_page_break()

    document.add_heading('Testing Best Practices', level=1)

    # Example usage: Loop through testing best practices questions
    for question in questionnaires['testing_best_practices']:
        response = ask_question(cache_name, question)
        document.add_paragraph(f'Question: {question}')
        document.add_paragraph(f'Answer: {response}')
        
    document.save(absolute_file_name)

    return { "code_review_file_name" : absolute_file_name }


# Node
def upload_file(state: GraphState):

    """ Upload code review file to the GCS bucket """

    file_name = state.get("code_review_file_name")
    repository_name = state.get("git_repository_name")
    release_name = state.get("git_release_name")

    code_review_file_name = file_name[file_name.rindex("/")+1:len(file_name)]

    upload_filepath = f"code_review/{repository_name}/{release_name}/{code_review_file_name}"

    blob = bucket.blob(upload_filepath)
    blob.upload_from_filename(file_name)

    print(f"File {file_name} uploaded to {upload_filepath}.")

    # Delete files on VM disk
    os.remove(file_name)

    return { "gcs_path_review_file" : upload_filepath }


